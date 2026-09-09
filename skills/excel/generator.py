from __future__ import annotations

from pathlib import Path
from typing import Any, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def create_excel(
    filename: str,
    sheet_name: str,
    headers: List[str],
    rows: List[List[Any]],
) -> str:
    """Crée un classeur Excel XLSX dans output/ et retourne son chemin absolu."""
    project_root = Path(__file__).resolve().parents[2]
    output_dir = project_root / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    filename = str(filename).strip()
    if not filename:
        raise ValueError("Le nom du fichier Excel est obligatoire.")
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"

    # Empêche l'écriture en dehors du dossier output/.
    safe_name = Path(filename).name
    if safe_name in {".", ".."}:
        raise ValueError("Nom de fichier Excel invalide.")
    output_path = output_dir / safe_name

    clean_sheet_name = str(sheet_name).strip() or "Feuille1"
    invalid_characters = set("[]:*?/\\")
    if any(char in invalid_characters for char in clean_sheet_name):
        raise ValueError("Le nom de la feuille Excel contient un caractère interdit.")
    if len(clean_sheet_name) > 31:
        raise ValueError("Le nom de la feuille Excel ne doit pas dépasser 31 caractères.")

    clean_headers = [str(header) for header in (headers or [])]
    if not clean_headers:
        raise ValueError("La liste des en-têtes Excel est obligatoire.")

    clean_rows = rows or []
    expected_columns = len(clean_headers)
    for index, row in enumerate(clean_rows, start=1):
        if not isinstance(row, (list, tuple)):
            raise ValueError(f"La ligne {index} doit être une liste de valeurs.")
        if len(row) != expected_columns:
            raise ValueError(
                f"La ligne {index} contient {len(row)} valeur(s), "
                f"mais {expected_columns} colonne(s) sont attendues."
            )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = clean_sheet_name

    worksheet.append(clean_headers)
    for row in clean_rows:
        worksheet.append(list(row))

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_index, cells in enumerate(worksheet.columns, start=1):
        max_length = max(
            (len(str(cell.value)) if cell.value is not None else 0 for cell in cells),
            default=0,
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 10),
            50,
        )

    workbook.save(output_path)
    return str(output_path.resolve())
