from __future__ import annotations

import os
import re
import warnings
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

# Constantes métier
STATUTS_FTNC = {"Non démarrées", "En cours"}
PROGRAMMES_FTNC = {
    "Alpha-Jet", "ATL2", "Falcon 2000", "Falcon 2000EX", "Falcon 900",
    "IA63", "Mirage 2000", "Mirage F1", "Rafale B/C", "Rafale Marine",
}
PROGRAMMES_SUIVI = {
    "ALPHA JET", "ATL2", "BREGUET", "FALCON 2000", "FALCON 2000EX",
    "FALCON 900", "IA63", "LEARJET 200", "LEARJET 85", "MIRAGE 2000",
    "MIRAGE 3", "MIRAGE F1", "RAFALE B/C", "RAFALE D", "RAFALE M",
}
PRIORITES = ["Urgent", "Important", "Moyen", "Minimum"]
FLAGS = {"Urgent": "🚩", "Important": "🟠", "Moyen": "🔵", "Minimum": "⚪", "Inconnue": "❔"}


def _get_ftnc_config() -> dict[str, Any]:
    """Charge la configuration du skill FTNC depuis config.json ou les variables d'environnement."""
    ftnc_cfg = {}
    try:
        from src.config.manager import load_config
        config = load_config()
        ftnc_cfg = config.get("skills", {}).get("ftnc", {})
    except (OSError, ValueError, ImportError):
        ftnc_cfg = {}

    env_ftnc = os.environ.get("FTNC_FICHIER_PLANNER")
    env_feuille_ftnc = os.environ.get("FTNC_FEUILLE_PLANNER")
    env_suivi = os.environ.get("FTNC_FICHIER_SUIVI_EURO")
    env_feuille_suivi = os.environ.get("FTNC_FEUILLE_SUIVI_EURO")

    fichier_ftnc = env_ftnc or ftnc_cfg.get("fichier_ftnc") or ""
    feuille_ftnc = env_feuille_ftnc or ftnc_cfg.get("feuille_ftnc") or "Données consolidées"
    fichier_suivi = env_suivi or ftnc_cfg.get("fichier_suivi_euro") or ""
    feuille_suivi = env_feuille_suivi or ftnc_cfg.get("feuille_suivi_euro") or "SUIVI"

    project_root = Path(__file__).resolve().parents[2]

    def _resolve(p_str: str) -> Path | None:
        if not p_str or not str(p_str).strip():
            return None
        p = Path(os.path.expandvars(os.path.expanduser(str(p_str).strip())))
        if not p.is_absolute():
            p = project_root / p
        return p

    return {
        "fichier_ftnc": _resolve(fichier_ftnc),
        "feuille_ftnc": str(feuille_ftnc).strip(),
        "fichier_suivi_euro": _resolve(fichier_suivi),
        "feuille_suivi_euro": str(feuille_suivi).strip(),
    }


def _texte(v: Any) -> str:
    if v is None or pd.isna(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _date(v: Any) -> str:
    if v is None or pd.isna(v):
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%y")
    texte = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(texte, fmt).replace(tzinfo=timezone.utc).strftime("%d/%m/%y")
        except ValueError:
            pass
    return texte


def _id10(v: Any) -> str | None:
    chiffres = "".join(re.findall(r"\d", _texte(v)))
    return chiffres[:10] if len(chiffres) >= 10 else None


def _normaliser_reference(v: Any) -> str:
    return "".join(c for c in _texte(v).casefold() if c.isalnum())


def _verifier(path: Path | None, role: str = "FTNC") -> Path:
    """Valide l'existence d'un fichier source Excel avec messages clairs."""
    if path is None or not str(path).strip():
        raise FileNotFoundError(
            f"Le fichier Excel {role} n'est pas configuré.\n"
            f"Veuillez définir 'skills.ftnc' dans config.json "
            f"ou la variable d'environnement FTNC_FICHIER_PLANNER / FTNC_FICHIER_SUIVI_EURO."
        )
    if not path.exists():
        raise FileNotFoundError(
            f"Le fichier Excel {role} est introuvable :\n{path}\n"
            f"Vérifiez le chemin renseigné dans config.json ('skills.ftnc')."
        )
    return path


def _lire_planner() -> pd.DataFrame:
    cfg = _get_ftnc_config()
    path = _verifier(cfg["fichier_ftnc"], "du planner FTNC (fichier_ftnc)")
    feuille = cfg["feuille_ftnc"]
    df = pd.read_excel(path, sheet_name=feuille, usecols="B,C,E,F", engine="openpyxl", dtype=str)
    if len(df.columns) != 4:
        raise ValueError(f"Impossible d'identifier les colonnes B, C, E et F de {path.name}.")
    df.columns = ["reference", "programme", "statut", "priorite"]
    for col in df.columns:
        df[col] = df[col].astype("string").str.strip()
    df = df[df["statut"].isin(STATUTS_FTNC) & df["programme"].isin(PROGRAMMES_FTNC)].copy()
    df = df.dropna(subset=["reference"])
    df["identifiant"] = df["reference"].apply(_id10)
    ordre = {p: i for i, p in enumerate(PRIORITES)}
    df["ordre"] = df["priorite"].map(ordre).fillna(999)
    return df.sort_values(["ordre", "reference"], kind="stable").drop(columns="ordre")


def _lire_suivi() -> list[dict[str, Any]]:
    cfg = _get_ftnc_config()
    path = _verifier(cfg["fichier_suivi_euro"], "du suivi €uro (fichier_suivi_euro)")
    feuille = cfg["feuille_suivi_euro"]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if feuille not in wb.sheetnames:
            raise ValueError(f'La feuille "{feuille}" est introuvable dans {path.name}.')
        ws = wb[feuille]
        autorises = {p.casefold() for p in PROGRAMMES_SUIVI}
        resultats = []
        for numero_ligne, ligne in enumerate(ws.iter_rows(min_row=2, max_col=18, values_only=True), start=2):
            numero, type_ftnc = ligne[0], ligne[3]
            ref_piece, designation, programme = ligne[5], ligne[6], ligne[7]
            quantite, description, pole = ligne[8], ligne[9], ligne[10]
            date_debut, statut = ligne[14], ligne[17]
            if _texte(statut).casefold() != "en cours" or _texte(pole).casefold() != "ppm":
                continue
            if _texte(programme).casefold() not in autorises:
                continue
            prefixe = "MWB" if _texte(type_ftnc).casefold() == "structure" else "VLB" if _texte(type_ftnc).casefold() == "carbone" else ""
            qte = _texte(quantite)
            try:
                pluriel = float(qte.replace(" ", "").replace(",", ".")) > 1
            except ValueError:
                pluriel = False
            resultats.append({
                "ligne_excel": numero_ligne,
                "numero": _texte(numero),
                "reference": f"{prefixe}{_texte(numero)}",
                "identifiant": _id10(numero),
                "programme": _texte(programme),
                "type": _texte(type_ftnc),
                "reference_piece": _texte(ref_piece),
                "designation": _texte(designation),
                "quantite": qte,
                "description": _texte(description),
                "date_debut": _date(date_debut),
                "statut": _texte(statut),
                "pole": _texte(pole),
                "resume": f"[{prefixe}{_texte(numero)}] {_texte(programme)} - {_texte(designation)} ({_texte(ref_piece)}) - {qte} {'pièces' if pluriel else 'pièce'}",
            })
        return resultats
    finally:
        wb.close()


def ftnc_liste() -> str:
    """Retourne la liste des FTNC et les nouvelles FTNC à ajouter au planner."""
    planner = _lire_planner()
    suivi = _lire_suivi()
    lignes = ["FTNC EN COURS", "=" * 40]
    for numero, (_, row) in enumerate(planner.iterrows(), start=1):
        priorite = _texte(row["priorite"]) or "Non renseignée"
        lignes.append(f"{numero}. {row['reference']} {FLAGS.get(priorite, FLAGS['Inconnue'])} | {row['programme']} | {row['statut']} | Priorité : {priorite}")
    if planner.empty:
        lignes.append("Aucune FTNC ne correspond aux critères demandés.")
    lignes.extend(["", f"Nombre de FTNC du planner : {len(planner)}"])

    occurrences = Counter(x for x in planner["identifiant"].tolist() if x)
    nouvelles = []
    for item in suivi:
        identifiant = item["identifiant"]
        if not identifiant:
            continue
        if occurrences[identifiant] > 0:
            occurrences[identifiant] -= 1
        else:
            nouvelles.append(item)
    if nouvelles:
        lignes.extend(["", "NOUVELLES FTNC À AJOUTER AU PLANNER", "-" * 40])
        for numero, item in enumerate(nouvelles, start=len(planner) + 1):
            lignes.extend([f"{numero}. {item['resume']}", item["description"] or "Description non renseignée.", f"Date de début : {item['date_debut'] or 'Non renseignée'}", ""])
    else:
        lignes.extend(["", "Aucune nouvelle FTNC à ajouter au planner."])
    lignes.append(f"Nombre de FTNC en cours dans le suivi €uro : {len(suivi)}")
    return "\n".join(lignes).strip()


def ftnc_details(reference: str) -> str:
    """Retourne les détails d'une FTNC trouvée par référence."""
    recherche = _normaliser_reference(reference)
    if not recherche:
        raise ValueError("La référence FTNC est obligatoire.")
    planner = _lire_planner()
    suivi = _lire_suivi()
    resultats: list[dict[str, Any]] = []
    for _, row in planner.iterrows():
        candidats = {_normaliser_reference(row["reference"]), _normaliser_reference(row["identifiant"])}
        if any(recherche == c or recherche in c for c in candidats if c):
            resultats.append({"source": "FTNC.xlsx / planner", **row.to_dict()})
    for item in suivi:
        candidats = {_normaliser_reference(item["reference"]), _normaliser_reference(item["numero"]), _normaliser_reference(item["identifiant"])}
        if any(recherche == c or recherche in c for c in candidats if c):
            resultats.append({"source": "Suivi des FTNC €uro.xlsx", **item})
    if not resultats:
        return f'Aucune FTNC trouvée pour la référence "{reference}".'

    lignes = [f'DÉTAILS DE LA FTNC "{reference}"', "=" * 40]
    for index, item in enumerate(resultats, start=1):
        if len(resultats) > 1:
            lignes.append(f"Correspondance {index}/{len(resultats)}")
        lignes.extend([
            f"Source : {item['source']}",
            f"Référence FTNC : {item.get('reference') or item.get('numero') or 'Non renseignée'}",
            f"Programme : {item.get('programme') or 'Non renseigné'}",
            f"Statut : {item.get('statut') or 'Non renseigné'}",
        ])
        if item["source"] == "FTNC.xlsx / planner":
            lignes.append(f"Priorité : {_texte(item.get('priorite')) or 'Non renseignée'}")
        else:
            lignes.extend([
                f"Type : {item.get('type') or 'Non renseigné'}",
                f"Pôle : {item.get('pole') or 'Non renseigné'}",
                f"Pièce : {item.get('designation') or 'Non renseignée'}",
                f"Référence pièce : {item.get('reference_piece') or 'Non renseignée'}",
                f"Quantité : {item.get('quantite') or 'Non renseignée'}",
                f"Date de début : {item.get('date_debut') or 'Non renseignée'}",
                f"Description : {item.get('description') or 'Non renseignée'}",
            ])
        lignes.append("")
    return "\n".join(lignes).strip()
