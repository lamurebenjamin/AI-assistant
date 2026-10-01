"""Assemblage et enregistrement du DOCX FTNC."""
import tempfile

from docx import Document
from docx.shared import Cm

from . import ftnc_formatting as formatting
from .ftnc_formatting import *
from .ftnc_images import add_extract_image
from .ftnc_layout import _value, add_title, apply_document_styles


def build_docx(output_path="document_sans_macro.docx", image_1_path=None, image_2_path=None, data=None):


    previous_temp_dir = formatting._TEMP_IMAGE_DIR
    temp_dir = tempfile.TemporaryDirectory(prefix="docx_generated_images_")
    formatting._TEMP_IMAGE_DIR = temp_dir

    try:
        document = Document()
        apply_document_styles(document)

        section = document.sections[0]
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)

        add_title(document, data=data)

        document.add_heading("DESCRIPTION DE LA NON-CONFORMITÉ", level=1)
        p_description = add_paragraph(document)
        add_run(p_description, _value(data, "Description", "Description de la non-conformité issue de CASQIT"), color=NOIR, italic=True, highlight=False)

        document.add_heading("ANALYSE", level=1)
        add_extract_image(document, image_1_path, _value(data, "PNR", "C20XXXXXX"), "XX", 1, caption_ref_highlight=False)
        add_extract_image(document, image_2_path, _value(data, "PNR", "AXXXX"), "XX", 2, caption_ref_highlight=False)

        document.add_heading("JUSTIFICATION", level=1)
        document.add_heading("FIT, FORM, FUNCTION", level=2)

        document.add_heading("FORM (ENVELOPPE)", level=3)
        add_instruction_paragraph(document, "Analyse (- matière / + Matière, Détection visuelle, …)")
        add_highlight_paragraph(document, "Justification Form")
        add_grey_conclusion(document, "Conclusion Form")

        document.add_heading("FUNCTION (FONCTION)", level=3)
        add_instruction_paragraph(document, "Analyse (description de la fonction)")
        add_highlight_paragraph(document, "Justification Function")
        add_grey_conclusion(document, "Conclusion Function")

        document.add_heading("FIT (ASSEMBLAGE)", level=3)
        add_instruction_paragraph(document, "Analyse (Chaîne de cote, règle métier/conception, …)")
        add_highlight_paragraph(document, "Justification Fit")
        add_grey_conclusion(document, "Conclusion Fit")

        document.add_heading("CONCLUSION FIT/FORM/FUNCTION", level=3)
        add_colored_macro_conclusion(document, "Conclusion FFF")

        add_page_break(document)

        document.add_heading("TENUE STRUCTURELLE", level=2)
        add_instruction_paragraph(document, "Tenue statique, durée de vie, fatigue, endurance")
        add_highlight_paragraph(document, "Justification Tenue structurelle")
        add_colored_macro_conclusion(document, "Conclusion Tenue Structurelle")

        document.add_heading("PERFORMANCE TRAITEMENT DE SURFACE", level=2)
        add_instruction_paragraph(document, "Protection contre la corrosion et traitement de surface")
        add_highlight_paragraph(document, "Justification Traitement de surface")
        add_colored_macro_conclusion(document, "Conclusion Traitement Surface")

        document.add_heading("RÉPARATION & MAINTENANCE", level=2)
        add_instruction_paragraph(document, "Capacité de réparation et procédures de maintenance")
        add_highlight_paragraph(document, "Justification Réparation / Maintenance")
        add_colored_macro_conclusion(document, "Conclusion Réparation Maintenance")

        add_page_break(document)
        document.add_heading("CONCLUSION", level=1)

        p = add_paragraph(document)
        add_run(p, "Suivant déclaration et analyse BE cette FTNC est classifiée : ")
        add_dropdown_run(
            p,
            control_name="Classification",
            placeholder="Sélectionner une classification",
            items=["MINEUR", "MAJEUR"],
            color=NOIR,
            bold=True,
        )

        p = add_paragraph(document)
        add_run(p, "Disposition pour les ", color=NOIR)
        add_run(p, _value(data, "Quantité", "XX"), color=NOIR, highlight=False)
        add_run(p, " pièces (n° OF / Série-Lot : ", color=NOIR)
        add_run(p, _value(data, "No. Série/Lot", "XXXX"), color=NOIR, highlight=False)
        add_run(p, ") : ", color=NOIR)
        add_dropdown_run(
            p,
            control_name="Disposition",
            placeholder="Sélectionner une disposition",
            items=["Acceptable en l’état", "Retouche", "Rebut"],
            color=NOIR,
            bold=True,
        )

        add_highlight_paragraph(document, "Synthèse à compléter.")

        document.save(output_path)
        return str(output_path)

    finally:
        formatting._TEMP_IMAGE_DIR = previous_temp_dir
        temp_dir.cleanup()
