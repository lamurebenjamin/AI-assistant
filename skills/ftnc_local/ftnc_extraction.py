"""Extraction des champs de la fiche NQ DOCX."""

import re

from docx import Document


def get_header_text(document):
    """
    Récupère tout le texte contenu dans les headers du document,
    y compris le texte présent dans les tableaux.
    """
    texts = []

    for section in document.sections:
        header = section.header

        for paragraph in header.paragraphs:
            if paragraph.text.strip():
                texts.append(paragraph.text.strip())

        for table in header.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        texts.append(cell_text)

    return "\n".join(texts)


def get_body_text(document):
    """
    Récupère tout le texte du corps du document,
    y compris le texte présent dans les tableaux.
    """
    texts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            texts.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    texts.append(cell_text)

    return "\n".join(texts)


def extract_value(text, label):
    """
    Extrait la valeur située après un libellé donné.
    La valeur est récupérée jusqu'à la fin de la ligne.
    """
    pattern = rf"{re.escape(label)}\s*(.*)"
    match = re.search(pattern, text, re.IGNORECASE)

    if match:
        return match.group(1).strip()

    return None


def extract_multiline_value(text, label, stop_labels=None):
    """
    Extrait une valeur pouvant être sur plusieurs lignes.
    La capture commence après le libellé et s'arrête avant le prochain libellé connu.
    """
    if stop_labels is None:
        stop_labels = []

    escaped_label = re.escape(label)

    if stop_labels:
        escaped_stop_labels = [re.escape(stop_label) for stop_label in stop_labels]
        stop_pattern = "|".join(escaped_stop_labels)
        pattern = rf"{escaped_label}\s*(.*?)(?=\n(?:{stop_pattern})|\Z)"
    else:
        pattern = rf"{escaped_label}\s*(.*)"

    match = re.search(pattern, text, re.IGNORECASE | re.DOTALL)

    if match:
        return match.group(1).strip()

    return None


def normalize_text(text):
    """
    Normalise le texte pour faciliter les comparaisons.
    """
    return re.sub(r"\s+", " ", text).strip().lower()


def extract_cell_below_label(document, label):
    """
    Recherche un libellé dans les tableaux du document,
    puis extrait le contenu de la cellule située juste en dessous.
    """
    label_normalized = normalize_text(label)

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell_text_normalized = normalize_text(cell.text)

                if label_normalized in cell_text_normalized:
                    next_row_index = row_index + 1

                    if next_row_index < len(table.rows):
                        next_row = table.rows[next_row_index]

                        if cell_index < len(next_row.cells):
                            value = next_row.cells[cell_index].text.strip()
                            return value if value else None

    return None


def extract_nq_data(docx_path):
    """
    Extrait les informations principales du fichier NQ.docx.
    """
    document = Document(docx_path)

    header_text = get_header_text(document)
    body_text = get_body_text(document)

    stop_labels = [
        "Programme Avion:",
        "Desc. PN:",
        "PN Commandé:",
        "Qté en dérogation:",
        "Description de la Non Conformité:",
        "No. Série/Lot Article",
        "No. Série/Lot"
    ]

    data = {
        "No. NC": extract_value(header_text, "No. NC:"),
        "Programme": extract_value(body_text, "Programme Avion:"),
        "Designation": extract_value(body_text, "Desc. PN:"),
        "PNR": extract_value(body_text, "PN Commandé:"),
        "Quantité": extract_value(body_text, "Qté en dérogation:"),
        "No. Série/Lot": extract_cell_below_label(document, "No. Série/Lot Article"),
        "Description": extract_multiline_value(
            body_text,
            "Description de la Non Conformité:",
            stop_labels=stop_labels
        )
    }

    return data

# =====================================================================
# Générateur FTNC issu du fichier FTNC.py, adapté pour recevoir les données NQ
