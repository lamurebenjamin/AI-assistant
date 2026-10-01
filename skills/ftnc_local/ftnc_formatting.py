"""Constantes et composants de mise en forme du document FTNC."""

from pathlib import Path

from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# =========================
# PARAMÈTRES GÉNÉRAUX
# =========================
FONT_SIZE_STEP = 1.0  # +1 pt sur toutes les tailles de police du document
PLACEHOLDER_IMAGE_WIDTH = 1800
PLACEHOLDER_IMAGE_HEIGHT = 780
PLACEHOLDER_FONT_MAX_SIZE = 96  # taille fortement augmentée pour les textes des 2 images générées
PLACEHOLDER_FONT_MIN_SIZE = 48
IMAGE_INSERT_WIDTH_INCHES = 6.2

BLEU_TITRE = "003366"
VERT = "196B24"
GRIS = "BFBFBF"
JAUNE = "D6D100"
JAUNE_SYMBOLE = "E3DE00"
ORANGE = "E97132"
ROUGE = "C00000"
NOIR = "000000"

POLICE_TEXTE = "Aptos (Body)"
POLICE_TITRE = "Aptos"

_TEMP_IMAGE_DIR = None


def adjusted_size(size):
    """Applique le cran supplémentaire demandé à toutes les tailles de police."""
    return float(size) + FONT_SIZE_STEP


def get_temp_image_path(filename):
    if _TEMP_IMAGE_DIR is None:
        return Path(filename)
    return Path(_TEMP_IMAGE_DIR.name) / filename


def set_rfonts(rPr, font_name=POLICE_TEXTE):
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)

    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)
    rFonts.set(qn("w:hint"), "default")

    for attr in ["w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:csTheme"]:
        attr_qn = qn(attr)
        if attr_qn in rFonts.attrib:
            del rFonts.attrib[attr_qn]


def set_document_default_font(document, font_name=POLICE_TEXTE):
    styles_element = document.styles.element
    doc_defaults = styles_element.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_element.insert(0, doc_defaults)

    rPr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rPr_default is None:
        rPr_default = OxmlElement("w:rPrDefault")
        doc_defaults.append(rPr_default)

    rPr = rPr_default.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        rPr_default.append(rPr)

    set_rfonts(rPr, font_name)


def set_run_font(run, font_name=POLICE_TEXTE, size=10, color=None, bold=False, italic=False, highlight=False):
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    set_rfonts(rPr, font_name)
    run.font.size = Pt(adjusted_size(size))
    run.bold = bold
    run.italic = italic

    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    return run


def add_run(paragraph, text, color=None, bold=False, italic=False, font=POLICE_TEXTE, size=10, highlight=False):
    return set_run_font(
        paragraph.add_run(text),
        font_name=font,
        size=size,
        color=color,
        bold=bold,
        italic=italic,
        highlight=highlight,
    )


def set_paragraph_spacing(paragraph, before=0, after=2, line_spacing=1.0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    return paragraph


def add_paragraph(document, text="", before=0, after=2, line_spacing=1.0):
    p = document.add_paragraph()
    set_paragraph_spacing(p, before, after, line_spacing)
    if text:
        add_run(p, text)
    return p


def add_page_break(document):
    """Sauts de page désactivés volontairement."""
    return


def set_paragraph_bottom_border(paragraph, color=BLEU_TITRE, size=18, space=8):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)

    bottom = pBdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pBdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)


def add_dropdown_run(
    paragraph,
    control_name,
    placeholder="Sélectionner une conclusion",
    items=None,
    font_name=POLICE_TEXTE,
    size=10,
    color=NOIR,
    italic=False,
    bold=False,
):
    """Ajoute un contrôle de contenu Word de type liste déroulante."""
    if items is None:
        items = [
            "✅ Pas d’impact",
            "⚠ Impact mineur (non notable)",
            "● Impact majeur (notable)",
            "⛔ Non acceptable",
        ]

    effective_size = adjusted_size(size)

    sdt = OxmlElement("w:sdt")
    sdtPr = OxmlElement("w:sdtPr")

    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), control_name)
    sdtPr.append(alias)

    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), control_name)
    sdtPr.append(tag)

    rPr_sdt = OxmlElement("w:rPr")
    set_rfonts(rPr_sdt, font_name)

    sz_sdt = OxmlElement("w:sz")
    sz_sdt.set(qn("w:val"), str(int(effective_size * 2)))
    rPr_sdt.append(sz_sdt)

    color_sdt = OxmlElement("w:color")
    color_sdt.set(qn("w:val"), color)
    rPr_sdt.append(color_sdt)

    if italic:
        rPr_sdt.append(OxmlElement("w:i"))
        rPr_sdt.append(OxmlElement("w:iCs"))
    if bold:
        rPr_sdt.append(OxmlElement("w:b"))
        rPr_sdt.append(OxmlElement("w:bCs"))

    sdtPr.append(rPr_sdt)

    dropdown = OxmlElement("w:dropDownList")
    for item in items:
        list_item = OxmlElement("w:listItem")
        list_item.set(qn("w:displayText"), item)
        list_item.set(qn("w:value"), item)
        dropdown.append(list_item)

    sdtPr.append(dropdown)
    sdt.append(sdtPr)

    sdtContent = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    set_rfonts(rPr, font_name)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(effective_size * 2)))
    rPr.append(sz)

    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    rPr.append(color_element)

    if italic:
        rPr.append(OxmlElement("w:i"))
        rPr.append(OxmlElement("w:iCs"))
    if bold:
        rPr.append(OxmlElement("w:b"))
        rPr.append(OxmlElement("w:bCs"))

    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(rPr)
    run.append(text)
    sdtContent.append(run)
    sdt.append(sdtContent)
    paragraph._p.append(sdt)
    return paragraph


def add_highlight_paragraph(document, text, italic=False):
    p = add_paragraph(document)
    add_run(p, text, color=NOIR, italic=italic, highlight=True)
    return p


def add_instruction_paragraph(document, text):
    p = add_paragraph(document, before=0, after=2)
    add_run(p, text, color=BLEU_TITRE, italic=True, size=8)
    return p


def add_mixed_highlight_paragraph(document, parts):
    p = add_paragraph(document)
    for text, highlight, color, bold in parts:
        add_run(p, text, color=color, bold=bold, highlight=highlight)
    return p


def add_grey_conclusion(document, conclusion_name):
    p = add_paragraph(document)
    add_run(p, "Conclusion :  ", color=GRIS, italic=True, size=10)
    add_dropdown_run(
        p,
        control_name=conclusion_name,
        placeholder="Sélectionner une conclusion",
        items=[
            "Pas d’impact",
            "Impact mineur (non notable)",
            "Impact majeur (notable)",
            "Non acceptable",
        ],
        color=GRIS,
        italic=True,
        bold=False,
    )
    return p


def add_colored_macro_conclusion(document, conclusion_name):
    p = add_paragraph(document)
    add_run(p, "Conclusion : ", color=NOIR, bold=True, size=10)
    add_dropdown_run(
        p,
        control_name=conclusion_name,
        placeholder="Sélectionner une conclusion",
        items=[
            "✅ Pas d’impact",
            "⚠ Impact mineur (non notable)",
            "● Impact majeur (notable)",
            "⛔ Non acceptable",
        ],
        color=NOIR,
        italic=False,
        bold=True,
    )
    return p
