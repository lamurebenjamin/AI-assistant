"""Images et légendes du document FTNC."""
import struct
import unicodedata
import zlib
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from .ftnc_formatting import add_paragraph, add_run, get_temp_image_path


def create_placeholder_image(path, watermark_text=None):
    """Crée une image PNG temporaire sans dépendance Pillow/PIL."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1200, 520
    white = (255, 255, 255)
    yellow = (245, 210, 0)
    pixels = [[white for _ in range(width)] for _ in range(height)]

    def put_px(x, y, color):
        if 0 <= x < width and 0 <= y < height:
            pixels[y][x] = color

    FONT = {
        "A":["01110","10001","10001","11111","10001","10001","10001"],"B":["11110","10001","10001","11110","10001","10001","11110"],
        "C":["01111","10000","10000","10000","10000","10000","01111"],"D":["11110","10001","10001","10001","10001","10001","11110"],
        "E":["11111","10000","10000","11110","10000","10000","11111"],"F":["11111","10000","10000","11110","10000","10000","10000"],
        "G":["01111","10000","10000","10011","10001","10001","01110"],"H":["10001","10001","10001","11111","10001","10001","10001"],
        "I":["11111","00100","00100","00100","00100","00100","11111"],"J":["00111","00010","00010","00010","10010","10010","01100"],
        "K":["10001","10010","10100","11000","10100","10010","10001"],"L":["10000","10000","10000","10000","10000","10000","11111"],
        "M":["10001","11011","10101","10101","10001","10001","10001"],"N":["10001","11001","10101","10011","10001","10001","10001"],
        "O":["01110","10001","10001","10001","10001","10001","01110"],"P":["11110","10001","10001","11110","10000","10000","10000"],
        "Q":["01110","10001","10001","10001","10101","10010","01101"],"R":["11110","10001","10001","11110","10100","10010","10001"],
        "S":["01111","10000","10000","01110","00001","00001","11110"],"T":["11111","00100","00100","00100","00100","00100","00100"],
        "U":["10001","10001","10001","10001","10001","10001","01110"],"V":["10001","10001","10001","10001","10001","01010","00100"],
        "W":["10001","10001","10001","10101","10101","10101","01010"],"X":["10001","10001","01010","00100","01010","10001","10001"],
        "Y":["10001","10001","01010","00100","00100","00100","00100"],"Z":["11111","00001","00010","00100","01000","10000","11111"],
        "0":["01110","10001","10011","10101","11001","10001","01110"],"1":["00100","01100","00100","00100","00100","00100","01110"],
        "2":["01110","10001","00001","00010","00100","01000","11111"],"3":["11110","00001","00001","01110","00001","00001","11110"],
        "4":["00010","00110","01010","10010","11111","00010","00010"],"5":["11111","10000","10000","11110","00001","00001","11110"],
        "6":["01110","10000","10000","11110","10001","10001","01110"],"7":["11111","00001","00010","00100","01000","01000","01000"],
        "8":["01110","10001","10001","01110","10001","10001","01110"],"9":["01110","10001","10001","01111","00001","00001","01110"],
        "-":["00000","00000","00000","11111","00000","00000","00000"],"'":["00100","00100","00000","00000","00000","00000","00000"]," ":["00000","00000","00000","00000","00000","00000","00000"],
    }

    def normalize_text(value):
        return "".join(c for c in unicodedata.normalize("NFD", value) if unicodedata.category(c) != "Mn").upper()

    def draw_char(ch, x, y, scale=4, color=yellow):
        pattern = FONT.get(ch, FONT[" "])
        for row, line_pattern in enumerate(pattern):
            for col, bit in enumerate(line_pattern):
                if bit == "1":
                    for dy in range(scale):
                        for dx in range(scale):
                            put_px(x + col * scale + dx, y + row * scale + dy, color)

    def draw_text_bold(value, x, y, scale=4, color=yellow):
        for i, ch in enumerate(normalize_text(value)):
            cx = x + i * 6 * scale
            draw_char(ch, cx, y, scale, color)
            draw_char(ch, cx + 1, y, scale, color)

    if watermark_text:
        normalized = normalize_text(watermark_text)
        if "ZONE NON CONFORME" in normalized:
            lines = ["EXTRAIT PLAN DE LA PIECE NON-CONFORME", "AVEC ZONE NON-CONFORME SURLIGNEE"]
        else:
            lines = ["EXTRAIT PLAN D'ENSEMBLE AVEC PIECE", "NON-CONFORME AFFICHEE EN JAUNE"]
        scale = 4
        line_h = 40
        y0 = (height - len(lines) * line_h) // 2
        for i, line_text in enumerate(lines):
            norm = normalize_text(line_text)
            x = max(10, (width - len(norm) * 6 * scale) // 2)
            draw_text_bold(norm, x, y0 + i * line_h, scale, yellow)

    raw = bytearray()
    for y in range(height):
        raw.append(0)
        for x in range(width):
            raw.extend(pixels[y][x])

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(bytes(raw), 9)) + chunk(b"IEND", b"")
    path.write_bytes(png)
    return str(path)


def ensure_image(image_path, default_name, watermark_text=None):
    """Utilise l’image fournie si elle existe, sinon crée une image temporaire."""
    if image_path and Path(image_path).exists():
        return str(image_path)
    default_path = get_temp_image_path(default_name)
    if watermark_text or not default_path.exists():
        create_placeholder_image(default_path, watermark_text=watermark_text)
    return str(default_path)


def add_drop_shadow_to_last_picture(run):
    """Ajoute une ombre portée autour de la dernière image insérée."""
    drawing = run._r.xpath(".//w:drawing")
    if not drawing:
        return
    pic_sp_pr = drawing[0].xpath(".//pic:spPr")
    if not pic_sp_pr:
        return
    sp_pr = pic_sp_pr[0]
    prst_geom = sp_pr.find(qn("a:prstGeom"))
    if prst_geom is None:
        prst_geom = OxmlElement("a:prstGeom")
        prst_geom.set("prst", "rect")
        av_lst = OxmlElement("a:avLst")
        prst_geom.append(av_lst)
        sp_pr.append(prst_geom)
    for old_effect in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(old_effect)
    effect_lst = OxmlElement("a:effectLst")
    outer = OxmlElement("a:outerShdw")
    outer.set("blurRad", str(23 * 12700))
    outer.set("dist", str(11 * 12700))
    outer.set("dir", str(45 * 60000))
    outer.set("algn", "ctr")
    outer.set("rotWithShape", "0")
    outer.set("sx", "100000")
    outer.set("sy", "100000")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", "333333")
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "65000")
    srgb.append(alpha)
    outer.append(srgb)
    effect_lst.append(outer)
    sp_pr.append(effect_lst)


def add_extract_image(document, image_path=None, caption_ref="AXXXX", caption_rev="XX", index=1, caption_ref_highlight=True):
    if index == 1:
        watermark = "Extrait plan d'ensemble avec pièce non-conforme affichée en jaune"
    elif index == 2:
        watermark = "EXTRAIT PLAN DE LA PIECE NON-CONFORME AVEC ZONE NON CONFORME SURLIGNEE"
    else:
        watermark = None
    image_file = ensure_image(image_path, f"extrait_{index}.png", watermark_text=watermark)
    p_img = add_paragraph(document, before=24, after=24)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    run.add_picture(image_file, width=Inches(5.8))
    add_drop_shadow_to_last_picture(run)
    p_caption = add_paragraph(document, before=0, after=0)
    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_caption, "Extrait ", italic=True)
    add_run(p_caption, caption_ref, italic=True, highlight=caption_ref_highlight)
    add_run(p_caption, " Rev.", italic=True)
    add_run(p_caption, caption_rev, italic=True, highlight=True)
