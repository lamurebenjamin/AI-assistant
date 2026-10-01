from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

from .ftnc_formatting import *


def set_style_font(style, font_name=POLICE_TEXTE, size=None, color=None, bold=True, italic=None, before=None, after=None):
    style.font.name = font_name
    rPr = style._element.get_or_add_rPr()
    set_rfonts(rPr, font_name)

    if size is not None:
        style.font.size = Pt(adjusted_size(size))
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if before is not None:
        style.paragraph_format.space_before = Pt(before)
    if after is not None:
        style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.0


def apply_document_styles(document):
    set_document_default_font(document, POLICE_TEXTE)
    styles = document.styles

    set_style_font(styles["Normal"], size=10, color=NOIR, bold=False, before=0, after=2)
    set_style_font(styles["Title"], font_name=POLICE_TITRE, size=19, color=BLEU_TITRE, bold=True, before=0, after=12)
    set_style_font(styles["Heading 1"], font_name=POLICE_TITRE, size=15, color=BLEU_TITRE, bold=True, before=18, after=12)
    set_style_font(styles["Heading 2"], font_name=POLICE_TITRE, size=13, color=BLEU_TITRE, bold=True, before=18, after=0)
    set_style_font(styles["Heading 3"], font_name=POLICE_TITRE, size=11, color=BLEU_TITRE, bold=True, before=12, after=0)


def _value(data, key, default=""):
    """Retourne une valeur exploitable pour le document, avec repli si champ absent."""
    if not data:
        return default
    value = data.get(key)
    if value is None or str(value).strip() == "":
        return default
    return str(value).strip()


def add_title(document, data=None):
    """Ajoute le titre principal en utilisant les données extraites de la NQ, sans surlignage."""
    no_nc = _value(data, "No. NC", "MWB2026XXXXXX")
    programme = _value(data, "Programme", "Programme")
    designation = _value(data, "Designation", "Désignation")
    pnr = _value(data, "PNR", "AXXXX")
    quantite = _value(data, "Quantité", "XX")

    p_title = document.add_paragraph(style="Title")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(18)

    add_run(p_title, f"[{no_nc}] ", color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19, highlight=False)
    add_run(p_title, programme, color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19, highlight=False)

    p_title.add_run().add_break(WD_BREAK.LINE)

    add_run(p_title, "Équipement", color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19)
    add_run(p_title, " - ", color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19)
    add_run(p_title, designation, color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19, highlight=False)
    add_run(p_title, " (", color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19)
    add_run(p_title, pnr, color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19, highlight=False)
    add_run(p_title, ") - ", color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19)
    add_run(p_title, quantite, color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19, highlight=False)
    add_run(p_title, " pièces", color=BLEU_TITRE, bold=True, font=POLICE_TITRE, size=19)

    set_paragraph_bottom_border(p_title)
