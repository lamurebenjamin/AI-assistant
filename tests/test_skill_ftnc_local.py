"""Tests du skill FTNC local avec bases temporaires, sans données réelles."""
import json
import sqlite3
import sys
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlparse
from urllib.request import url2pathname

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from skills.ftnc_local import generator as g


class FtncLocalTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        root = Path(self.tmp.name)
        self.root = root
        pdm = root / 'PNR.db'
        with closing(sqlite3.connect(pdm)) as conn, conn:
            conn.executescript('CREATE TABLE articles(ref TEXT PRIMARY KEY, designation TEXT NOT NULL, issue TEXT, code_equivalence TEXT, "group" TEXT); CREATE TABLE bom(id INTEGER PRIMARY KEY, root_id INTEGER, parent_id INTEGER, parent_article_ref TEXT, article_ref TEXT NOT NULL, quantity REAL NOT NULL);')
            conn.executemany('INSERT INTO articles VALUES (?,?,?,?,?)', [('ROOT','Racine','A',None,None),('PROG','Programme','B',None,None),('EQUIP','Equipement','C',None,None),('REF-1','Piece','D','EQ','GR')])
            conn.executemany('INSERT INTO bom VALUES (?,?,?,?,?,?)', [(1,1,None,None,'ROOT',1),(2,1,1,'ROOT','PROG',1),(3,1,2,'PROG','EQUIP',1),(4,1,3,'EQUIP','REF-1',1)])
        with closing(sqlite3.connect(root / 'FTNC.db')) as conn, conn:
            conn.execute('CREATE TABLE FTNC(cle TEXT PRIMARY KEY, programme TEXT, equipment TEXT, designation TEXT, pnr TEXT, type TEXT, description TEXT, quantite TEXT, responsable TEXT, disposition TEXT, reponse TEXT, date_ouverture TEXT, date_fermeture TEXT, status TEXT)')
            conn.execute('INSERT INTO FTNC(cle,pnr,description) VALUES (?,?,?)', ('NC1','ref-1','exemple'))
        (root/'config.json').write_text(json.dumps({'skills':{'ftnc_local':{'pdm_database_path':'PNR.db','ftnc_database_path':'FTNC.db'}}}), encoding='utf-8')
        patcher=patch.object(g,'PROJECT_ROOT',root)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_recherche_article(self):
        result=g.rechercher_article_ftnc('REF-1')
        self.assertEqual(result['issue'],'D')
        self.assertEqual(result['programmes'],['PROG - EQUIP (REF-1)'])
        self.assertFalse(g.rechercher_article_ftnc('INCONNU')['trouve'])

    def test_ftnc_et_suggestions(self):
        self.assertEqual(g.rechercher_ftnc_existantes('REF-1')['count'],1)
        self.assertEqual(g.rechercher_ftnc_existantes('AUTRE')['count'],0)
        self.assertEqual(g.suggerer_references_ftnc('REF1')['suggestions'][0]['reference'],'REF-1')

    def test_fixture_connections_are_closed_before_cleanup(self):
        # Sur Windows, une base temporaire encore ouverte bloque sa suppression.
        for name in ('PNR.db', 'FTNC.db'):
            path = self.root / name
            renamed = self.root / (name + '.renamed')
            path.rename(renamed)
            renamed.rename(path)
        self.assertEqual(g.rechercher_ftnc_existantes('REF-1')['count'], 1)
        path = self.root / 'FTNC.db'
        renamed = self.root / 'FTNC.db.renamed'
        path.rename(renamed)
        renamed.rename(path)

    def test_validation(self):
        with self.assertRaises(ValueError): g.rechercher_article_ftnc('')
        with self.assertRaises(ValueError): g.suggerer_references_ftnc('REF', 0)
        with self.assertRaises(ValueError): g.generer_justification_ftnc('{}')
        with self.assertRaises(ValueError): g.extraire_fiche_ftnc('absent.docx')

    def test_extraction_et_generation_reelles(self):
        from docx import Document
        doc=Document()
        header=doc.sections[0].header
        header.add_paragraph('No. NC: NC-TEST-001')
        for line in ('Programme Avion: Programme Test', 'Desc. PN: Pièce Test',
                     'PN Commandé: REF-1', 'Qté en dérogation: 2',
                     'Description de la Non Conformité: Défaut de test'):
            doc.add_paragraph(line)
        path=self.root/'nq.docx'
        doc.save(path)
        extracted=g.extraire_fiche_ftnc(str(path))
        self.assertEqual(extracted['PNR'],'REF-1')
        self.assertEqual(extracted['No. NC'],'NC-TEST-001')
        self.assertEqual(extracted['Description'],'Défaut de test')
        extracted['Type FTNC']='Dimensionnel'
        with patch('skills.ftnc_local.FTNC_Justification_Generator.build_docm', side_effect=RuntimeError('Word indisponible')):
            result=g.generer_justification_ftnc(json.dumps(extracted,ensure_ascii=False))
        self.assertFalse(result['macro_enabled'])
        self.assertTrue(result['warning'])
        out=Path(result['fichier'])
        self.assertTrue(out.is_file())
        self.assertEqual(out.suffix,'.docx')
        generated=Document(out)
        self.assertIn('REF-1',generated.paragraphs[0].text)
        self.assertIn('NC-TEST-001',generated.paragraphs[0].text)

    def test_modules_under_quality_limit_and_public_facade(self):
        from skills.ftnc_local import FTNC_Justification_Generator as facade
        for name in ('extract_nq_data', 'build_output_filename', 'build_docm', 'build_docx'):
            self.assertTrue(callable(getattr(facade, name)))
        for module in (ROOT / 'skills' / 'ftnc_local').glob('*.py'):
            self.assertLessEqual(len(module.read_text(encoding='utf-8').splitlines()), 500, module.name)

    def _nq_for_orchestrator(self, reference='REF-1'):
        from docx import Document
        doc = Document()
        doc.sections[0].header.add_paragraph('No. NC: NC-001')
        for text in ('Programme Avion: Programme', 'Desc. PN: Pièce',
                     f'PN Commandé: {reference}', 'Qté en dérogation: 2',
                     'Description de la Non Conformité: Écart constaté'):
            doc.add_paragraph(text)
        path = self.root / 'fiche-nq.docx'
        doc.save(path)
        return path

    def test_orchestrateur_questions_puis_generation(self):
        path = self._nq_for_orchestrator()
        # La façade publique retourne une réponse épurée (statut + question)
        first = g.traiter_fiche_ftnc(str(path))
        self.assertEqual(first['statut'], 'questions')
        self.assertIn('question', first)
        self.assertNotIn('generation', first)
        # Les données internes sont accessibles via le point d'entrée détaillé
        detail = g._traiter_fiche_ftnc_detail(str(path))
        self.assertEqual(detail['statut'], 'questions')
        self.assertEqual(detail['article']['issue'], 'D')
        self.assertEqual(detail['ftnc_existantes']['count'], 1)
        self.assertEqual(len(detail['etapes']), 5)
        self.assertNotIn('generation', detail)
        corrections = json.dumps({'No. Série/Lot': 'LOT-1'})
        pending = g.traiter_fiche_ftnc(str(path), reference_confirmee='REF-1',
                                        type_ftnc='Dimensionnel', corrections_json=corrections)
        self.assertEqual(pending['statut'], 'questions')
        self.assertNotIn('generation', pending)
        with patch('skills.ftnc_local.FTNC_Justification_Generator.build_docm', side_effect=RuntimeError('Word absent')):
            done = g.traiter_fiche_ftnc(str(path), reference_confirmee='REF-1',
                                         type_ftnc='Dimensionnel', corrections_json=corrections,
                                         confirmer_generation=True)
        self.assertEqual(done['statut'], 'termine')
        self.assertIn('lien', done)
        self.assertTrue(Path(url2pathname(urlparse(done['fichier_url']).path)).is_file() if done.get('fichier_url') else True)

    def test_orchestrateur_ref_absente_et_liens(self):
        path = self._nq_for_orchestrator('INCONNU')
        result = g.traiter_fiche_ftnc(path.as_uri(), reference_confirmee='INCONNU',
                                      type_ftnc='Dimensionnel', confirmer_generation=True)
        self.assertEqual(result['statut'], 'questions')
        self.assertIn('suggestions', result)
        self.assertNotIn('generation', result)
        with self.assertRaises(ValueError):
            g.traiter_fiche_ftnc('https://example.org/nq.docx')
        with self.assertRaises(ValueError):
            g.traiter_fiche_ftnc(str(path), corrections_json='{"PNR":"REF-1"}')

    def test_schema_wrapper(self):
        try:
            from skills.ftnc_local.skill import FtncLocalSkill
        except ModuleNotFoundError as exc:
            if exc.name in ('core','core.mcp_compat'):
                self.skipTest('Infrastructure core.mcp_compat absente du dossier de test isolé')
            raise
        tools=FtncLocalSkill().get_tools()
        self.assertEqual(len(tools),1)
        self.assertEqual(tools[0]['name'],'traiter_fiche_ftnc')
        self.assertTrue(all(t['parameters']['additionalProperties'] is False for t in tools))

if __name__ == '__main__':
    unittest.main()
